"""Extend the retained Ass6, preserving its literature and eleven-section structure."""
import argparse
import copy
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT/'ass6_output/Residential_Intelligence_Computer_Project_Proposal_READY_2026-10-18.docx'
PRIVATE = Path.home()/'SeminarLab/preboard/document_review/revision'

def read(name):
    return json.loads((ROOT/'reports/preboard'/name).read_text(encoding='utf-8'))

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--final',action='store_true'); args=ap.parse_args()
    ocr=read('single_pass_ocr.json'); llm=read('single_pass_llm.json'); pc=read('pc_transport_profile.json')
    if args.final and (not llm['complete'] or pc['run_id']!='pc-delivery-profile-engineering-v3' or pc['failures']):
        raise ValueError('Final release requires completed LLM matrix and corrected-parser PC profile')
    doc=Document(SOURCE); old=list(doc.paragraphs)
    diagram=ROOT/'ass6_assets/system_architecture_preboard.png'
    blip=doc.inline_shapes[0]._inline.xpath('.//a:blip')[0]
    doc.part.related_parts[blip.get(qn('r:embed'))]._blob=diagram.read_bytes()
    doc.inline_shapes[0].height=int(doc.inline_shapes[0].width*820/1200)
    old[6].insert_paragraph_before('หมายเลขโครงงาน SCS3-4  กลุ่มสอบ SCS3')
    assert old[35].text.startswith('4.6') and old[91].text.startswith('6 ')
    changes=[]
    def replace(i,text):
        p=old[i]; original=p.text
        rpr=copy.deepcopy(p.runs[0]._r.rPr) if p.runs else None
        p.clear(); run=p.add_run(text)
        if rpr is not None: run._r.insert(0,rpr)
        changes.append({'source_paragraph':i,'before':original,'after':text})

    replace(37,'Cai และคณะเสนอ Edge-LLM เป็นกรอบการทำงานร่วมกันสำหรับให้บริการโมเดลภาษาขนาดใหญ่ในระบบ edge computing งานนี้ใช้เป็นฐานเปรียบเทียบการแบ่งภาระประมวลผลระหว่างองค์ประกอบของระบบ กับแนวทางประมวลผลหลักภายในบ้านของโครงงาน โดยยังไม่นำผลสมรรถนะของบทความมาอ้างแทนอุปกรณ์และข้อมูลของโครงงาน [18]')
    replace(23,'รายงานทบทวนวรรณกรรมเดิมเปรียบเทียบงานกลุ่ม quantization/hardware โดยชี้ว่าผลจาก Raspberry Pi, Jetson และ GPU ระดับศูนย์ข้อมูลถ่ายโอนมาใช้กับเครื่องเป้าหมายโดยตรงไม่ได้ งานของ Giedra และคณะพิจารณาความแม่นยำและทรัพยากรตามสภาพแวดล้อมที่ทดลอง ส่วน ShadowNPU ใช้ sparse attention ร่วมกับ NPU และ CPU/GPU เพื่อลดภาระจากส่วน attention ที่ไวต่อ quantization [2, 12, 17] โครงงานจึงบันทึกรุ่นโมเดล ระดับบีบอัด และเครื่องที่ทดสอบคู่กับผล JSON เสมอ การออกแบบแบบ ShadowNPU ยังเป็นแนวทางศึกษา มิใช่ผลการพัฒนาของโครงงาน')
    replace(14,old[14].text.replace('ผลนำร่องล่าสุด','ผลนำร่องระยะแรก')+' ผลเพิ่มเติมใช้ข้อความอ้างอิงหนึ่งรอบครบ 59 ภาพ เปรียบเทียบ OCR บนข้อความผ่านเกณฑ์ 24 ภาพ ทดลองรูปแบบ JSON ด้วย Local LLM และเชื่อมเดโมถึง Home Assistant บน PC โดยแสดงรายละเอียด ตัวหาร และข้อจำกัดในหัวข้อ 5.9–5.12 [23]')
    replace(29,'สถาปัตยกรรมเดิมระบุ Home Assistant, MQTT, การพยากรณ์แบบ Adaptive Holt-Winters และ dashboard เป็นปลายทางที่ต้องศึกษาความเป็นไปได้ [12, 15, 16] ระยะนี้เชื่อมผลจากภาพจริงไปยัง Home Assistant ในสภาพแวดล้อมทดลองแล้ว โดยบันทึกแหล่งที่มาและสถานะการยืนยัน การวัดทางเทคนิคใช้การยืนยันอัตโนมัติและแยกจากความถูกต้องของเนื้อหา ส่วน burn rate ต้องมีอนุกรมเวลาของปริมาณคงเหลือหรือการใช้จริงเพียงพอก่อน จึงยังไม่เป็นผลส่งมอบที่รับรองจากภาพใบเสร็จชุดปัจจุบัน [23]')
    replace(56,'ชั้นปลายน้ำแยกผลการรับส่งข้อความออกจากความถูกต้องของรายการสินค้า ปัจจุบันมีหลักฐานเชื่อมข้อมูลจากภาพจริงผ่าน MQTT ไปยัง Home Assistant โดยใช้การยืนยันอัตโนมัติในสภาพแวดล้อมทดลองและตรวจค่าที่ปลายทาง [23] ส่วนการพยากรณ์อัตราการใช้ยังไม่ได้ทดลอง เพราะยังไม่มีข้อมูลคงเหลือหรือการใช้ตามเวลาเพียงพอ ขอบเขตนี้รักษาแนวคิด Home Assistant และ Adaptive Holt-Winters จากรายงานเดิม [12, 15, 16]')
    replace(40,'Yin และคณะพัฒนา ShadowNPU โดยใช้ sparse attention และการประสาน NPU กับ CPU/GPU เพื่อลดภาระที่ต้องย้อนกลับไปคำนวณนอก NPU เนื่องจากความไวต่อ quantization งานนี้ช่วยชี้ว่าข้อจำกัดของตัวดำเนินการและการจัดตารางคำนวณมีผลต่อประสิทธิภาพ โครงงานใช้เป็นกรอบอภิปรายความเข้ากันได้ของ runtime โดยยังไม่ได้ติดตั้งเทคนิคดังกล่าวหรือทดสอบบน RK3588 [17]')
    replace(43,'Giedra และคณะประเมินโมเดลที่บีบอัดหลายขนาด โดยวัดความถูกต้องบน PC และวัดประสิทธิภาพการประมวลผลบน Jetson Orin Nano Super การระบุแพลตฟอร์มและค่าคอนฟิกแยกกันมีประโยชน์ต่อการออกแบบการทดลองของโครงงาน แต่ชุดทดสอบของบทความไม่ใช่การสกัด JSON จากใบเสร็จ จึงต้องวัดโครงสร้างและค่าฟิลด์เฉพาะงานเพิ่มเติม [2]')
    replace(59,'ภาพที่ 1 ขยายสถาปัตยกรรมสามชั้นจากรายงานทบทวนวรรณกรรมเดิมให้เห็น e-Senses ของบ้าน ชั้นรับรู้ใช้ e-Eye รับภาพใบซื้อสินค้า ตรวจกรอบด้วย YOLO และอ่านข้อความด้วย OCR ชั้นประมวลผลใช้ Local LLM สกัด JSON และตรวจ schema ก่อนให้ผู้ใช้ตรวจยืนยัน ชั้นประยุกต์ใช้ส่งข้อมูลผ่าน MQTT ไปยัง Home Assistant ต้นแบบบน PC เชื่อมกระบวนการนี้ได้แล้วในสภาพแวดล้อมทดลอง ส่วน IP Camera, e-Nose และ e-Skin ยังเป็นส่วนขยายตามแผน [12, 23]')
    replace(62,'ผลนำร่องเดิมประกอบด้วย YOLOv8n บน PC การตรวจ EasyOCR แบบ spot-check และการแปลง ONNX เป็น RKNN เป้าหมาย RK3588 แบบไม่ quantize ซึ่งรัน inference ผ่าน SDK บน PC ได้หนึ่งภาพ ภาพส่งเพิ่ม 6 ภาพมีตัวหนังสือหมุนข้าง เมื่อหมุนภาพและกรอบอ้างอิงให้ตรงโดยไม่ฝึกเพิ่ม mAP@0.5 เปลี่ยนจาก 0.237 เป็น 0.714 ส่วนการหมุน 7 ภาพที่อ่านตรงในชุดแรกไป 90 องศาทำให้คะแนนเปลี่ยนจาก 0.995 เป็น 0.492 ผลนี้แสดงความไวต่อการหมุนเฉพาะชุดเล็ก ผลทดลองเพิ่มเติมด้าน OCR, Local LLM และ MQTT รายงานในหัวข้อ 5.9–5.11 โดยยังไม่มีผลวัดบนบอร์ด [4, 9, 23]')
    replace(67,'การตรวจชื่อสินค้าเบื้องต้นใน 59 ภาพเคยแบ่งเป็นถอดชื่อครบครั้งแรก 12 ภาพ ถอดไม่ครบ 45 ภาพ และไม่มีบรรทัดชื่อสินค้า 2 ภาพ ตัวเลขนี้เป็นผลตรวจเฉพาะชื่อสินค้าในระยะก่อนหน้า ส่วนข้อมูลล่าสุดเป็นข้อความทุกบรรทัดที่ผู้จัดทำบันทึกหนึ่งรอบ 59 ภาพจากชุด 59 ภาพ จึงเป็นคนละขอบเขตการตรวจและใช้แทนกันไม่ได้ ยังไม่มีสัดส่วนภาษาไทย/อังกฤษของชื่อสินค้าที่รับรองครบทั้งชุด [4, 23]')
    replace(68,'ชุดข้อมูล 59 ภาพ มีข้อความอ้างอิงจากการตรวจหนึ่งรอบ 59 ภาพ โดยนายณัชพล เพ็งพลเป็นผู้บันทึก เก็บประวัติการแก้ไขและสถานะข้อความไม่ครบไว้ตามจริง ไม่มีภาพค้างบันทึกข้อความแล้ว การตรวจรอบนี้เน้นข้อความ จึงไม่รับรองช่อง JSON ที่เว้นว่างหรือมีค่าเริ่มต้นว่าเป็นข้อมูลไม่ปรากฏจริง schema กำหนด transaction_date, merchant, items[{name, quantity, unit, unit_price, line_total}], grand_total และ currency ส่วนสถานะอ่านไม่ออกเก็บในบันทึกประกอบ การวัดความถูกต้องของฟิลด์ยังค้างเพราะไม่มีเฉลย JSON ที่ตรวจครบ และไม่มีข้อกำหนดให้ผู้จัดทำกลับไปกรอกรอบสองในระยะนี้ [23]')
    replace(70,'การตรวจจับใช้กรอบ receipt_region ตามนิยามของชุดข้อมูล รายงาน mAP@0.5, mAP@0.5:0.95, precision, recall และสัดส่วนกรอบคะแนนสูงสุดที่ IoU≥0.5 แยกชุดและจำนวนภาพ การทดลอง OCR เปรียบเทียบ EasyOCR กับ Tesseract บนพิกเซลภาพ crop เดียวกัน ใช้ CER จากข้อความหนึ่งรอบหลังตัดเฉพาะป้ายอ้างอิง [cite: N] แล้วปรับ Unicode NFC และช่องว่างตามกฎที่บันทึกไว้ ผลนี้เป็นการเปรียบเทียบเชิงพรรณนา การวัดวันที่ ยอดรวม และรายการสินค้าแยกต่างหากยังต้องมีเฉลยฟิลด์ที่ตรวจครบ [6–8, 23]')
    replace(71,'การทดลอง Local LLM เปรียบเทียบ Qwen3 ขนาด 0.6B และ 1.7B ที่ Q8_0 และ Q4_K_M พร้อมเปิด/ปิด grammar ใช้ข้อความที่ผู้จัดทำบันทึกและข้อความ Tesseract ของภาพเดียวกัน ภายใต้ prompt, schema, seed และขีดจำกัดโทเคนคงที่ ตัวชี้วัดที่วัดได้ในระยะนี้คือ parse validity, schema validity, เวลา และ RSS ของกระบวนการโมเดล ส่วน exact match รายฟิลด์ ความคลาดเคลื่อนจำนวนเงิน และความครบของรายการยังไม่มีคะแนน เพราะไม่มีเฉลย JSON ที่ตรวจครบ การรันซ้ำใช้ศึกษาความแปรผันของเวลา ไม่เพิ่มจำนวนภาพหรือธุรกรรมอิสระ [23]')
    replace(72,old[72].text.replace('ปัจจุบันตัวให้คะแนนผ่านการทดสอบด้วยข้อมูลสมมติ ยังไม่มีผลประเมินโมเดล JSON จริง','ตัวให้คะแนนผ่านการทดสอบกรณี JSON ผิดรูปแบบและการจับคู่รายการด้วยข้อมูลทดสอบ ส่วนการรันโมเดลจริงมีผลตรวจ parse/schema และทรัพยากรแล้ว แต่ยังไม่ใช้ตัวให้คะแนนเนื้อหาจนกว่าจะมีเฉลยฟิลด์ครบ'))
    replace(79,'ต้นแบบบน PC ส่งผลจากภาพจริงผ่าน YOLO, OCR และ Local LLM ไปยังหน้าตรวจยืนยัน ก่อนส่ง MQTT และแสดงยอดรวมกับจำนวนรายการใน Home Assistant การทดสอบการส่งใช้การยืนยันอัตโนมัติแทนมนุษย์เฉพาะในสภาพแวดล้อมทดลอง และตรวจว่าค่าที่ปลายทางตรงกับ JSON ที่ส่ง จึงเป็นหลักฐานการเชื่อมต่อ ไม่ใช่คะแนนความถูกต้องของสินค้าและไม่ใช่การศึกษาผู้ใช้ ตามฐานรายงานเดิมจะพิจารณา Adaptive Holt-Winters เมื่อมีข้อมูลคงเหลือหรือการใช้หลายช่วงเวลา ใบเสร็จบอกการซื้อเข้า จึงยังไม่มีผล burn rate [12, 15, 16, 23]')
    replace(82,'ข้อมูลนำเข้าแต่ละภาพมีรหัสภาพและค่า SHA-256 สำหรับตรวจย้อนกลับ ระบบบันทึกกรอบเอกสาร คะแนนความเชื่อมั่น การหมุนภาพ และรุ่นโมเดลก่อน crop เข้าสู่ OCR จากนั้นเก็บข้อความดิบ คำขอและคำตอบ LLM พร้อม JSON ที่แยกได้โดยไม่เขียนทับต้นฉบับ ข้อมูลยืนยันและข้อความ MQTT ผูกกับรหัสการประมวลผลและ event_id ส่วนการรับรองกลุ่มธุรกรรมซ้ำและการผูกค่าทุกฟิลด์กับบริเวณข้อความยังไม่ครบ จึงไม่อ้างว่าทำ provenance ระดับฟิลด์สำเร็จทั้งหมด')
    replace(83,'ก่อนส่ง MQTT ระบบตรวจ JSON ด้วย parser ที่ปฏิเสธคีย์ซ้ำและค่าตัวเลขที่ไม่ใช่จำนวนจำกัด จากนั้นตรวจ schema และรับการยืนยันผ่านหน้าตรวจ ข้อกำหนดอนุญาต null ในฟิลด์ที่ไม่มีข้อมูลและไม่บังคับว่าต้องมีรายการสินค้า การผ่าน schema จึงไม่ได้พิสูจน์ความถูกต้องหรือความครบของเนื้อหา ข้อความส่งออกมี event_id และสถานะยืนยันเพื่อป้องกันการตีความข้อมูลที่ยังรอตรวจเป็นรายการยืนยันแล้ว')
    replace(88,'ข้อจำกัดของการทดลองคือข้อความอ้างอิงผ่านผู้บันทึกหนึ่งคนหนึ่งรอบและอาจมีการจัดรูปแบบใหม่ มีผล OCR ให้ดูระหว่างกรอก จึงไม่ใช่เฉลยอิสระที่ปิดบังคำตอบโมเดล ภาพใกล้ซ้ำและกลุ่มธุรกรรมยังไม่รับรองครบ การทดลองใช้ภาพเดิมและยังไม่มีชุดภาพต่างฉากใหม่ที่แยกจากการพัฒนา ผลจึงรายงานเชิงพรรณนาพร้อมตัวหาร ไม่สร้างช่วงความเชื่อมั่นที่สมมติว่าภาพทุกใบอิสระ และไม่สรุปสมรรถนะทั่วไปของระบบ [23]')
    replace(90,'แผนส่งมอบยังสืบจาก roadmap เจ็ดระยะในรายงานเดิม ปัจจุบันมีชุดภาพและข้อความหนึ่งรอบ ผล OCR และผลรูปแบบ JSON บน PC รวมถึงหลักฐานเชื่อม MQTT/Home Assistant แล้ว งานที่ยังค้างคือการรับรองกรอบและกลุ่มธุรกรรม เฉลย JSON สำหรับคะแนนเนื้อหา และการทดสอบภาพต่างฉากใหม่ ก่อนสรุปประสิทธิภาพปลายทาง ส่วนการวิเคราะห์การใช้สินค้าและการวัดบน RK3588 เป็นระยะต่อเนื่องเมื่อมีข้อมูลและอุปกรณ์ [12, 23]')
    replace(93,'ชุดข้อมูลยังมีข้อจำกัดด้านข้อความไม่ครบ กลุ่มธุรกรรมซ้ำ และกรอบที่ยังไม่รับรองทั้งหมด รายงาน CER จึงใช้เพียงภาพที่ผ่านเกณฑ์คัดข้อความของการเปรียบเทียบหนึ่งรอบ ส่วนผล JSON รายงานเฉพาะการอ่านรูปแบบและการตรง schema ยังไม่มี field accuracy หรือ end-to-end semantic accuracy อีกทั้งยังไม่มีชุดภาพต่างฉากใหม่สำหรับทดสอบสุดท้าย ผล NPU และพลังงานจริงยังค้าง ข้อจำกัดเหล่านี้ทำให้งานวิจัยก่อนมีบอร์ดทั้งห้าข้อยังไม่ถือว่าทดลองครบ [23]')
    replace(102,'Implementation Roadmap เจ็ดระยะจากรายงานเดิมยังเป็นแผนแม่บท ได้แก่ ทบทวนวรรณกรรมและสถาปัตยกรรม ต้นแบบ YOLO/OCR การสกัด JSON การเชื่อม MQTT/Home Assistant การประเมินข้อมูลนำร่อง จริยธรรมเมื่อศึกษากับผู้ใช้ และการวัด trade-off ครบระบบ ระยะ PC มีผลเชื่อมระบบและวัดรูปแบบ/ทรัพยากรแล้ว แต่ผลเนื้อหาและการทดสอบต่างฉากใหม่ยังค้าง ตารางที่ 8 แสดงแผนต่อเนื่องนับจากการรับรองเค้าโครง [12, 23]')
    # Retain the original roadmap's schedule while updating its evidence requirements.
    for p in doc.paragraphs:
        if p.text.startswith('ตารางที่ 4 '):
            p.text=p.text.replace('ตารางที่ 4','ตารางที่ 8',1)
        if p.text.startswith('9  แผนและระยะเวลาดำเนินการ'):
            p.paragraph_format.page_break_before=False
    doc.tables[3].cell(2,1).text='ใช้ข้อความหนึ่งรอบ และกำหนดเฉลยฟิลด์สำหรับระยะถัดไป'
    doc.tables[3].cell(2,2).text='ข้อความ 59 ภาพ; เฉลย JSON ยังค้าง'
    doc.tables[3].cell(1,2).text='กรอบและกลุ่มธุรกรรมที่รับรองแล้ว'
    doc.tables[3].cell(5,2).text='ข้อจำกัดและแผนทดลอง\nบนบอร์ดจริง'
    # Insert added results before section 6, using original paragraph styles.
    anchor=old[91]
    def para(text,style='Normal'):
        p=anchor.insert_paragraph_before(text,style)
        p.paragraph_format.keep_together=True
        if style.startswith('Heading'):
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.left_indent=Inches(.25)
        return p
    def table(number,title,headers,rows,widths):
        cap=para(f'ตารางที่ {number}  {title}')
        cap.paragraph_format.keep_with_next=True
        cap.paragraph_format.space_after=Pt(14)
        t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.autofit=False
        for i,h in enumerate(headers): t.rows[0].cells[i].text=h
        for row in rows:
            for c,v in zip(t.add_row().cells,row): c.text=str(v)
        for row in t.rows:
            for c,w in zip(row.cells,widths): c.width=Inches(w)
        anchor._p.addprevious(t._tbl)
        para('')
        return t
    para('5.9  ผลเปรียบเทียบ OCR กับข้อความอ้างอิงหนึ่งรอบ','Heading 2')
    para('ประเมินข้อความที่บันทึก 59 ภาพโดยคงข้อมูลดิบทุกภาพไว้ พบ 24 ภาพผ่านเกณฑ์คัดเข้า และ 35 ภาพไม่ใช้คำนวณ CER เต็มข้อความ เนื่องจากมีข้อความอ่านไม่ชัด ตารางที่จัดใหม่ หมายเหตุที่ไม่ใช่ข้อความในภาพ ตัวเลขย่อ หรือขอบเขตข้อความไม่ครบ หนึ่งภาพอาจมีหลายเหตุผล เกณฑ์อัตโนมัติและการตรวจภาพเฉพาะกรณียังไม่รับรองว่าข้อความ 24 ภาพถูกต้องครบทุกตัวอักษร [23]')
    para('CER คำนวณจากจำนวนการแทนที่ ลบ และแทรกตัวอักษรในการจัดแนวข้อความ หารด้วยจำนวนตัวอักษรอ้างอิง โดยรวมตัวเศษและตัวหารทุกภาพสำหรับ micro CER หลังตัดเฉพาะป้ายอ้างอิง [cite: N] จากข้อความทั้งสองฝั่ง แล้วปรับ Unicode NFC และช่องว่างตามโค้ดประเมิน ค่าต่ำหมายถึงผิดพลาดน้อยและอาจเกินร้อยละ 100 ได้ ตัวหารชุดหลักมี 28,011 ตัวอักษร ทั้งสองเครื่องมือใช้ภาพ crop เดียวกันที่ตรวจค่าแฮชแล้ว ผล EasyOCR และ Tesseract จึงเทียบกันบนข้อความอ้างอิงชุดเดียวกัน [23]')
    summary=ocr['summaries']
    rows=[]
    for cohort,label in [('all_eligible','ผ่านเกณฑ์คัดเข้า'),('without_declared_ocr_assistance','ไม่แจ้งใช้ OCR ช่วย'),('declared_ocr_assistance','แจ้งใช้ OCR ช่วย')]:
        e=summary[cohort]['easyocr']; t=summary[cohort]['tesseract']
        rows.append([label,e['scorable_documents'],f"{e['reference_characters']:,}",f"{100*e['micro_cer']:.2f}%",f"{100*t['micro_cer']:.2f}%"])
    table(4,'ความผิดพลาดตัวอักษรเชิงพรรณนา (micro CER)', ['ข้อความอ้างอิง','ภาพ','ตัวอักษร','EasyOCR','Tesseract'],rows,[1.85,.45,.8,1.05,1.05])
    para('EasyOCR มี CER ต่ำกว่า Tesseract บนชุดที่คัดนี้ แต่ทั้งสองยังมีความผิดพลาดสูง การไม่แจ้งใช้ OCR ช่วยเป็นเพียงข้อมูลที่ผู้บันทึกระบุ ไม่ยืนยันความเป็นอิสระของเฉลย จึงยังไม่เลือกเครื่องมือจากผลนี้เพื่ออ้างประสิทธิภาพทั่วไป ผล Tesseract ใช้การรันใหม่ทั้ง 59 ภาพหลังแก้ตัวอ่าน TSV ที่เคยตีความเครื่องหมายอัญประกาศผิด ผลจากตัวอ่านรุ่นเดิมเก็บเป็นประวัติและไม่ใช้คำนวณคะแนนนี้ [23]')
    para('5.10  ผล Local LLM ด้านรูปแบบและต้นทุนการประมวลผล','Heading 2')
    para(f"ชุดทดลองเลือก 6 ภาพที่มี crop เอกสารเดียวจากลำดับความยาวข้อความก่อนดูผลโมเดล โดยตรึงการเลือกตั้งแต่มีข้อความบันทึก 55 ภาพ ข้อความของหกภาพนี้ไม่เปลี่ยนหลังเติมครบ 59 ภาพ การเลือกนี้ไม่ใช่การสุ่มตัวแทนทั้งชุด ใช้ข้อความผู้บันทึกและ Tesseract ของแต่ละภาพ เปรียบเทียบ Qwen3 0.6B/1.7B ที่ Q8_0/Q4_K_M เปิดและปิด grammar และรันซ้ำ 3 รอบ รวมแผน 288 ครั้ง บันทึกแล้ว {llm['attempts']} ครั้ง สถานะ {'ครบตามเมทริกซ์' if llm['complete'] else 'กำลังดำเนินการ ยังไม่ครบเมทริกซ์'} การรันซ้ำไม่ใช่ภาพเพิ่มอีกชุด [23]")
    para('ใช้ PC Intel Core Ultra 7 155H, RAM ติดตั้งประมาณ 16 GB และ Windows 11 รัน llama.cpp b11714 บน CPU จำนวน 6 threads, context 8,192 tokens, output limit 2,048 tokens, temperature 0 และ seed 20261018 ปิด reasoning และ prompt cache พร้อม warm-up ก่อนจับเวลาแต่ละการโหลดโมเดล Q4_K_M แปลงต่อจาก Q8_0 ของผู้เผยแพร่ จึงไม่ใช่การเปรียบเทียบ quantization ที่เริ่มจาก BF16 เดียวกัน โมเดลและอินพุตมี SHA-256 พร้อมคำขอและคำตอบดิบ ผลเวลาอาจได้รับอิทธิพลจากงานเบื้องหลังบน PC [23]')
    for number,source,label in [(5,'owner_text','ข้อความผู้บันทึก'),(6,'tesseract','ข้อความ Tesseract')]:
        rows=[]
        for model in ['qwen3-0.6B-Q8_0','qwen3-0.6B-Q4_K_M','qwen3-1.7B-Q8_0','qwen3-1.7B-Q4_K_M']:
            for grammar in [False,True]:
                g=next((r for r in llm['groups'] if r['model']==model and r['source']==source and r['grammar']==grammar),None)
                if g:
                    rows.append([model.replace('qwen3-',''), 'เปิด' if grammar else 'ปิด',f"{g['parse_valid']}/{g['attempts']}",f"{g['schema_valid']}/{g['attempts']}",f"{g['median_seconds']:.2f}",f"{g['max_model_process_tree_rss_bytes']/2**30:.2f}"])
                else: rows.append([model.replace('qwen3-',''),'เปิด' if grammar else 'ปิด','ยังไม่รัน','ยังไม่รัน','—','—'])
        table(number,f'ผล LLM จาก{label}', ['โมเดล','บังคับ','Parse','Schema','วินาที¹','GiB²'], rows,[1.4,.65,.75,.75,.85,.8])
        para('¹ มัธยฐานเวลาต่อคำขอ รวมครั้งล้มเหลวที่บันทึก  ² RSS สูงสุดที่สุ่มวัดของกระบวนการโมเดล แต่ละเงื่อนไขครบเมื่อมี 18 ครั้งจาก 6 ภาพ × 3 รอบ ค่า Parse/Schema เป็นจำนวนผ่านหารจำนวนครั้งที่รันจริง')
    para('การตรง schema ยืนยันเพียงโครงสร้างและชนิดข้อมูล แม้ผลลัพธ์ใช้ null หรือรายการว่างก็อาจผ่านได้ ยังไม่มีคะแนนความถูกต้องวันที่ ยอดรวม หรือรายการสินค้า และไม่สรุปว่าโมเดลหรือ grammar ใดแม่นยำกว่าในเชิงเนื้อหาจากตารางนี้ [23]')
    if (ROOT/'reports/preboard/completed_run_audit.json').exists():
        audit=read('completed_run_audit.json'); termination=audit['llm_status_counts']
        para(f"การตรวจสถานะคำตอบพบ truncated จากเพดาน 2,048 tokens จำนวน {termination.get('truncated',0)} ครั้ง และคำขอล้มเหลว {termination.get('request_failed',0)} ครั้ง ทุกครั้งยังรวมในตัวหาร ตาราง parse/schema ตรวจจากข้อความดิบที่โมเดลคืนมาโดยไม่มีการซ่อมคำตอบ จึงต้องอ่านควบคู่กับสถานะการจบคำตอบ [23]")
        if audit['grammar_truncated']:
            para(f"ในกลุ่มที่เปิด grammar พบคำตอบถูกตัด {audit['grammar_truncated']} ครั้ง จึงต้องตรวจว่าคำตอบจบครบก่อนใช้งาน การบังคับรูปแบบช่วยจำกัดโทเคนที่สร้างได้ แต่ไม่ได้ยกเลิกเพดานความยาว และยังไม่ยืนยันความถูกต้องเชิงเนื้อหา [23]")
    para('5.11  เดโมครบกระบวนการและทรัพยากรบน PC','Heading 2')
    if pc.get('configuration'):
        cfg=pc['configuration']
        para(f"ชุดวัด PC ใช้ YOLOv8n ตรวจกรอบ เครื่องมือ OCR {cfg['engine']} และโมเดล {cfg['model_id']} โดย {'เปิด' if cfg['grammar'] else 'ปิด'} grammar กำหนดชุดนี้เพื่อวัดการเชื่อมต่อและทรัพยากร หมุนภาพตามทิศที่บันทึกไว้โดยผู้จัดทำ จึงยังไม่ใช่การตรวจทิศภาพอัตโนมัติ [23]")
    para('ต้นแบบเชื่อมภาพจริงกับการตรวจกรอบ OCR Local LLM หน้าตรวจยืนยัน MQTT และ Home Assistant แล้ว หลักฐานทดสอบการเชื่อมต่อใช้ 3 ภาพ ตรวจ event_id ยอดรวม และจำนวนรายการที่ปลายทาง รวมทั้งกรณี broker หยุดแล้วส่งซ้ำ การยืนยันซ้ำ และ JSON ไม่ผ่าน schema การยืนยันในชุดทดสอบนี้เป็นคำสั่งอัตโนมัติเฉพาะสภาพแวดล้อมทดลอง ไม่ใช่การตรวจเนื้อหาโดยผู้ใช้ [23]')
    s=pc['summaries']; rows=[]
    for key,label,scale in [('automated_window_to_ha_observed_seconds','ภาพถึงสังเกตค่าใน HA (วินาที)',1),('publish_to_ha_observed_seconds','เริ่มส่งถึงสังเกตค่าใน HA (วินาที)',1),('windows_full_host_rss_bytes','Windows host RSS (GiB)',2**30),('ha_linux_service_rss_bytes','Home Assistant RSS (MiB)',2**20),('broker_linux_service_rss_bytes','MQTT broker RSS (MiB)',2**20)]:
        v=s[key]; rows.append([label,v['samples'],f"{v['median']/scale:.3f}",f"{v['min']/scale:.3f}–{v['max']/scale:.3f}"])
    para(f"ผลวัดอ้างอิงการรัน {pc['run_id']} ใช้ {pc['source_photos']} ภาพ × 3 รอบ รวม {pc['attempts']} ครั้ง ล้มเหลว {pc['failures']} ครั้ง โหลดโมเดลค้างในหน่วยความจำและตัด warm-up ออกจากสถิติ {'ผลนี้ยังเป็นรุ่นก่อนแก้ TSV และจะรันทวนก่อนส่งฉบับสุดท้าย' if pc['run_id']!='pc-delivery-profile-engineering-v3' else 'ใช้ตัวอ่าน Tesseract TSV รุ่นแก้ไขแล้ว'} [23]")
    table(7,'เวลาและหน่วยความจำของเดโม PC', ['ตัวชี้วัด','ครั้ง','มัธยฐาน','ต่ำสุด–สูงสุด'],rows,[2.65,.45,1,1.1])
    if pc.get('stage_summaries'):
        stages=pc['stage_summaries']
        labels=[('detection_and_preprocessing','ตรวจกรอบและเตรียมภาพ'),('ocr_windows_client','OCR รวมการเรียกจาก Windows'),('llm_server','Local LLM'),('confirmation_persist_surrogate','บันทึกการยืนยันอัตโนมัติ'),('mqtt_client','MQTT client จนรับ PUBACK และปิดการเชื่อมต่อ')]
        details=[]
        for key,label in labels:
            v=stages[key]
            details.append(f"{label} {v['wall_seconds']['median']:.3f} วินาที และ RSS สูงสุดที่สุ่มวัด {v['sampled_peak_rss_bytes']['max']/2**30:.3f} GiB")
        para('เมื่อแยกขั้นตอน มัธยฐานเวลาและหน่วยความจำที่วัดได้คือ '+ '; '.join(details)+' โดย RSS ของการตรวจกรอบและ OCR ฝั่ง Windows รวมโมเดลที่โหลดค้างในกระบวนการ ไม่ใช่หน่วยความจำที่เพิ่มขึ้นเฉพาะขั้นตอนนั้น [23]')
        v=stages['ocr_linux_child']
        para(f"เฉพาะโปรเซส Tesseract ใน Linux มีมัธยฐานเวลา {v['wall_seconds']['median']:.3f} วินาที และ peak RSS สูงสุด {v['child_peak_rss_bytes']['max']/2**20:.3f} MiB วัดด้วย RUSAGE_CHILDREN แยกจากค่า Windows client จึงไม่รวมค่าทั้งสองเป็นหน่วยความจำสุทธิ [23]")
    para('เวลาเริ่มส่งถึงสังเกตค่า Home Assistant รวมรอบ HTTP polling จึงเป็นขอบเขตเวลาสังเกต ไม่ใช่เวลา processing ภายใน HA โดยตรง เวลาเต็มกระบวนการแทนช่วงผู้ใช้ตรวจด้วยการยืนยันอัตโนมัติและไม่รวมเวลาตรวจของมนุษย์ RSS ฝั่ง Windows และบริการ Linux วัดคนละขอบเขต จึงไม่บวกเป็น RAM รวมระบบ การรันซ้ำ 3 ภาพเป็นผลเชิงวิศวกรรม ไม่ใช่คะแนนทดสอบทั่วไปหรือหลักฐานว่าใช้ได้บน NPU [23]')
    para('5.12  สถานะงานก่อนมีบอร์ดและสิ่งที่ยังค้าง','Heading 2')
    para('(1) มีภาพ 59 ภาพและข้อความหนึ่งรอบ 59 ภาพ แต่กรอบและกลุ่มธุรกรรมยังไม่รับรองครบ (2) มี CER เชิงพรรณนา 24 ภาพ (3) มีผลรัน LLM ด้านรูปแบบและทรัพยากร (4) มีเดโมส่งข้อมูลถึง Home Assistant และการตรวจค่าปลายทางในสภาพแวดล้อมทดลอง (5) มีการวัดเวลาและ RAM บน PC การวัดความถูกต้องเชิงเนื้อหาของฟิลด์ JSON และการทดสอบภาพต่างฉากใหม่ระบุสถานะว่ายังไม่ได้ทดลอง เนื่องจากยังไม่มีเฉลยฟิลด์ที่ตรวจครบและชุดภาพใหม่ตามลำดับ ส่วนคะแนน NPU และพลังงานต้องรออุปกรณ์จริง [23]')
    # New local evidence reference, immediately before signature area.
    ref=next(p for p in old if p.text.startswith('ลงชื่อ')).insert_paragraph_before('[23] ณัชพล เพ็งพล และ ธนบดี เริ่มรักสกุล. (2569). ผลทดลองก่อนมีบอร์ด: ข้อความอ้างอิงหนึ่งรอบ OCR Local LLM และการเชื่อม MQTT/Home Assistant. ชุดหลักฐานพร้อมค่าแฮชและข้อจำกัด. https://github.com/natchapolp-glitch/Capstone-Project/tree/pc-preboard-implementation/reports/preboard')
    ref.paragraph_format.left_indent=Inches(.25); ref.paragraph_format.first_line_indent=Inches(-.25)
    # Make all short tables indivisible, including their existing captions.
    for t in doc.tables:
        t.autofit=False
        widths=[c.width for c in t.rows[0].cells]
        for column,width in zip(t.columns,widths):column.width=width
        tw=t._tbl.tblPr.find(qn('w:tblW'))
        if tw is not None:
            tw.set(qn('w:type'),'dxa'); tw.set(qn('w:w'),str(round(sum(widths)/635)))
        for ri,row in enumerate(t.rows):
            pr=row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
            if ri==0: pr.append(OxmlElement('w:tblHeader'))
            for ci,c in enumerate(row.cells):
                c.width=widths[ci]
                tcpr=c._tc.get_or_add_tcPr()
                old_borders=tcpr.find(qn('w:tcBorders'))
                if old_borders is not None:tcpr.remove(old_borders)
                borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    border=OxmlElement('w:'+edge)
                    for key,value in [('val','single'),('sz','4'),('color','000000')]:border.set(qn('w:'+key),value)
                    borders.append(border)
                tcpr.append(borders)
                for pi,p in enumerate(c.paragraphs):
                    p.paragraph_format.keep_with_next=ri<len(t.rows)-1
                    p.paragraph_format.keep_together=True
                    for r in p.runs:
                        r.font.name='TH Sarabun New'; r.font.size=Pt(14)
                        rpr=r._r.get_or_add_rPr(); rf=rpr.rFonts
                        if rf is not None:
                            for key in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+key),'TH Sarabun New')
                        sz=OxmlElement('w:szCs'); sz.set(qn('w:val'),'28'); rpr.append(sz)
                        if ri==0:
                            r.bold=True
                            b=OxmlElement('w:bCs'); b.set(qn('w:val'),'1'); rpr.append(b)
    for p in doc.paragraphs:
        if re.match(r'^\[\d+\]',p.text):
            p.paragraph_format.keep_together=True
        if p.text.startswith('ตารางที่ '):
            value=p.text
            p.clear(); prefix,title=value.split('  ',1)
            r=p.add_run(prefix); r.bold=True
            b=OxmlElement('w:bCs'); b.set(qn('w:val'),'1'); r._r.get_or_add_rPr().append(b)
            p.add_run('  '+title)
            p.paragraph_format.keep_with_next=True
    PRIVATE.mkdir(parents=True,exist_ok=True)
    suffix='REVIEWED' if args.final else 'DRAFT'
    target=PRIVATE/f'Ass6_{suffix}.docx'
    doc.save(target)
    (PRIVATE/'ass6_changes.json').write_text(json.dumps({'source':str(SOURCE),'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'changes':changes,'new_sections':['5.9','5.10','5.11','5.12'],'output':str(target)},ensure_ascii=False,indent=2),encoding='utf-8')
    print(target)

if __name__=='__main__': main()
