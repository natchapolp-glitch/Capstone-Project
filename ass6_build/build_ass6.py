from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "ass6_output" / "Residential_Intelligence_Computer_Project_Proposal_READY_2026-10-17.docx"
ARCH = ROOT / "ass6_assets" / "system_architecture.png"
OUT.parent.mkdir(exist_ok=True)
FONT = "TH Sarabun New"

def font_style(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.rFonts
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn(f"w:{key}"), FONT)

def set_run_font(run, bold=None, size=None):
    run.font.name = FONT
    if bold is not None:
        run.bold = bold
    if size is not None:
        run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.rFonts
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn(f"w:{key}"), FONT)

def p(text="", style=None, align=None, keep=False):
    para = doc.add_paragraph(style=style)
    if text:
        set_run_font(para.add_run(text))
    if align is not None:
        para.alignment = align
    para.paragraph_format.keep_together = keep
    return para

def heading(number, title):
    para = p(f"{number}  {title}", "Heading 1")
    para.paragraph_format.keep_with_next = True
    return para

def subheading(number, title):
    para = p(f"{number}  {title}", "Heading 2")
    para.paragraph_format.keep_with_next = True
    return para

def bullet(text):
    para = p(text, "List Bullet")
    para.paragraph_format.left_indent = Inches(0.25)
    para.paragraph_format.first_line_indent = Inches(-0.13)
    return para

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.style = "Table Grid"
    if widths:
        for row in t.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        pr = c.paragraphs[0]
        pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(pr.add_run(h), bold=True, size=13)
        tcpr = c._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "DCE6EE")
        tcpr.append(shd)
    trpr = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trpr.append(repeat)
    for ri, data in enumerate(rows):
        row = t.add_row()
        for i, val in enumerate(data):
            c = row.cells[i]
            c.text = ""
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = c.paragraphs[0]
            pr.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(pr.add_run(str(val)), size=13)
            if ri % 2:
                tcpr = c._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "F4F7F9")
                tcpr.append(shd)
    for row in t.rows:
        for c in row.cells:
            tcpr = c._tc.get_or_add_tcPr()
            mar = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                e = OxmlElement(f"w:{edge}")
                e.set(qn("w:w"), "80")
                e.set(qn("w:type"), "dxa")
                mar.append(e)
            tcpr.append(mar)
            for par in c.paragraphs:
                par.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def caption(kind, number, label):
    para = doc.add_paragraph()
    set_run_font(para.add_run(f"{kind}ที่ {number}  "), bold=True, size=14)
    set_run_font(para.add_run(label), size=14)
    para.paragraph_format.keep_with_next = True
    para.paragraph_format.space_before = Pt(9)
    para.paragraph_format.space_after = Pt(8)
    return para

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(1.5)
sec.left_margin = Inches(1.5)
sec.right_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.header_distance = Inches(0.5)
sec.different_first_page_header_footer = True

normal = doc.styles["Normal"]
font_style(normal, 14)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.widow_control = True
for name, size, before, after in (("Title", 18, 0, 12), ("Heading 1", 16, 13, 6), ("Heading 2", 14, 9, 4)):
    s = doc.styles[name]
    font_style(s, size, True)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.keep_together = True
    if name == "Title":
        ppr = s.element.get_or_add_pPr()
        for border in ppr.findall(qn("w:pBdr")):
            ppr.remove(border)
font_style(doc.styles["List Bullet"], 14)
doc.styles["List Bullet"].paragraph_format.space_after = Pt(2)

# Page number at the top right, omitted on the first page.
page_header = sec.header.paragraphs[0]
page_header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
page_header._p.append(field)

title = p("การเสนอเค้าโครงโครงงานคอมพิวเตอร์", "Title", WD_ALIGN_PARAGRAPH.CENTER)
p("ภาควิชาวิทยาการคอมพิวเตอร์  คณะวิทยาศาสตร์  มหาวิทยาลัยขอนแก่น", align=WD_ALIGN_PARAGRAPH.CENTER)
p("นายณัชพล เพ็งพล  รหัสนักศึกษา 673380267-4")
p("นายธนบดี เริ่มรักสกุล  รหัสนักศึกษา 673380270-5")
p("นักศึกษาระดับปริญญาตรี  หลักสูตรวิทยาศาสตรบัณฑิต  สาขาวิชาวิทยาการคอมพิวเตอร์")
p("อาจารย์ที่ปรึกษาโครงงาน  ผศ.บุญทรัพย์ ไวคำ")

heading("1", "ชื่อหัวข้อโครงงาน")
p("ภาษาไทย  การพัฒนาระบบประสาทส่วนกลางอัจฉริยะในที่พักอาศัยเพื่อสกัดและวิเคราะห์ข้อมูลห่วงโซ่อุปทานในครัวเรือนด้วยปัญญาประดิษฐ์ประมวลผลเฉพาะถิ่น")
p("ภาษาอังกฤษ  DEVELOPMENT OF A LOCALIZED RESIDENTIAL INTELLIGENCE ENGINE FOR HOUSEHOLD SUPPLY-CHAIN ANALYTICS USING EDGE COMPUTING AND HOME NETWORK INFRASTRUCTURE")

heading("2", "หลักการและเหตุผล")
p("โครงงานเสนอระบบประสาทส่วนกลางอัจฉริยะในที่พักอาศัยที่รับข้อมูลจากประสาทสัมผัสดิจิทัลหลายชนิด หรือ e-Senses แล้วแปลงเป็นข้อมูลที่ตรวจสอบและใช้วิเคราะห์ได้ e-Eye รับภาพเอกสารซื้อสินค้า e-Nose อาจรับข้อมูลก๊าซหรือกลิ่น และ e-Skin อาจรับน้ำหนักหรือแรงสัมผัส ระยะแรกเลือก e-Eye เพื่อทำการทดลองที่มีข้อมูลจริง ส่วนโหนดอื่นยังเป็นส่วนขยายในกรอบระบบ ไม่ใช่ผลทดลองที่สำเร็จแล้ว [12]")
p("ข้อมูลซื้อสินค้าในครัวเรือนอยู่ในใบเสร็จที่มีรูปแบบ ภาษา และคุณภาพภาพถ่ายต่างกัน การแปลงภาพเป็นรายการสินค้าแบบมีโครงสร้างต้องผ่านการตรวจจับเอกสาร การรู้จำข้อความ และการสกัดฟิลด์ ความผิดพลาดของขั้นต้นอาจส่งต่อถึงผล JSON และการวิเคราะห์ปลายน้ำ จึงต้องวัดทั้งความแม่นยำและทรัพยากรที่ใช้ในแต่ละชั้น")
p("รายงานทบทวนวรรณกรรม 32 หน้าของโครงงานสังเคราะห์บทความวิจัย 12 ฉบับในสามมิติ คือ quantization และฮาร์ดแวร์ การสร้างผลลัพธ์แบบมีโครงสร้าง และบริบทบ้านอัจฉริยะ/IoT งานที่ทบทวนให้แนวทางวัด trade-off แต่ยังไม่ตอบโดยตรงว่าขั้นตอนเหล่านี้ทำงานร่วมกันได้ดีเพียงใดกับใบซื้อสินค้าจริงภาษาไทย/อังกฤษภายใต้ทรัพยากรจำกัด ช่องว่างนี้กล่าวเฉพาะชุดวรรณกรรมที่ทบทวน ไม่ใช่ข้อสรุปว่าไม่มีงานอื่นศึกษาเรื่องเดียวกัน [1–3, 12]")
p("ผลนำร่องล่าสุดรับภาพ 124 ไฟล์ เหลือภาพไม่ซ้ำระดับไฟล์ 113 ภาพ และคัดภาพผู้สมัครใบซื้อสินค้ามีรายการ 59 ภาพ โมเดล YOLOv8n บน PC ได้ mAP@0.5 เท่ากับ 0.928 ในภาพตรวจ 10 ภาพจากชุดแรก และ 0.237 ในภาพส่งเพิ่ม 6 ภาพซึ่งกันไว้ทดสอบโมเดลเดิมโดยไม่ฝึกเพิ่ม ภาพชุดหลังต่างฉากด้วย จึงยังแยกสาเหตุของคะแนนที่ต่างกันไม่ได้ ต้องเพิ่มภาพและตรวจกรอบอ้างอิงก่อนสรุปประสิทธิภาพของระบบ [4]")

heading("3", "วัตถุประสงค์ของโครงงาน")
bullet("ออกแบบกรอบสถาปัตยกรรม e-Senses และกระบวนการสามชั้น ได้แก่ การรับรู้และสกัดข้อมูล การประมวลผลเฉพาะถิ่น และการใช้งานในระบบบ้านอัจฉริยะ โดยใช้ e-Eye เป็นกรณีศึกษาระยะแรก")
bullet("พัฒนาต้นแบบแปลงภาพใบซื้อสินค้าเป็นข้อมูล JSON ที่ตรวจสอบย้อนหลังได้ และวัดความผิดพลาดของการตรวจจับ OCR และการสกัดข้อมูลทั้งรายขั้นและปลายทาง")
bullet("เปรียบเทียบผลของโมเดลขนาดเล็ก ระดับ quantization และการบังคับรูปแบบ JSON ต่อความถูกต้อง เวลา และหน่วยความจำภายใต้ทรัพยากรจำกัด โดยเริ่มบน PC และตรวจการส่งออกโมเดลสำหรับ RK3588")
bullet("ออกแบบการส่งข้อมูลที่ผ่านการตรวจแล้วสู่ระบบบ้านอัจฉริยะ และกำหนดข้อมูลเพิ่มเติมที่จำเป็นต่อการประเมินปริมาณคงเหลือหรืออัตราการใช้สินค้า ก่อนทดลองบนอุปกรณ์จริงเมื่อมีความพร้อม")

heading("4", "ทฤษฎีและผลงานวิจัยที่เกี่ยวข้อง")
subheading("4.1", "การประมวลผลเฉพาะถิ่นและ trade-off")
p("การลด precision ของโมเดลช่วยลดภาระทรัพยากร แต่ต้องประเมินความแม่นยำร่วมกับเวลาและพลังงาน ไม่อาจเลือกจากขนาดไฟล์เพียงอย่างเดียว Husom และคณะศึกษาการอนุมาน quantized LLM บนอุปกรณ์ edge ขณะที่ Giedra และคณะเสนอการประเมินหลายตัวชี้วัดบนระบบที่มีทรัพยากรจำกัด [1, 2] งาน ShadowNPU ศึกษาการแบ่งงานระหว่างหน่วยประมวลผลบนอุปกรณ์ [17] ในโครงงานนี้การเปรียบเทียบจะยึดงานสกัดใบซื้อสินค้าไทย/อังกฤษแทน benchmark ทั่วไป")
subheading("4.2", "การสร้างผลลัพธ์แบบมีโครงสร้าง")
p("SGLang ศึกษาการเร่งการทำงานของโปรแกรมภาษาแบบมีโครงสร้าง [3] และ llama.cpp รองรับ grammar/JSON schema บางส่วน [5] โครงงานจะเปรียบเทียบการสร้าง JSON แบบเปิดกับแบบบังคับ schema โดยวัดทั้งความถูกต้องของโครงสร้างและค่าฟิลด์ เพราะ JSON ที่ parse ได้ยังอาจมีรายการสินค้าหรือราคาไม่ตรงหลักฐาน")
subheading("4.3", "การเชื่อมระบบบ้านอัจฉริยะและข้อมูลปลายน้ำ")
p("รายงาน Assignment 4 สังเคราะห์งานบ้านอัจฉริยะสามแนวทาง: ตัวแทน LLM สำหรับควบคุมบ้าน [14], การพยากรณ์ข้อมูล IoT แบบอนุกรมเวลา [15] และเกตเวย์ ESP32 สำหรับการเชื่อมต่อ [16] โครงงานนี้จะส่งเฉพาะข้อมูลรายการซื้อที่ผ่านการตรวจแล้วผ่าน MQTT และแยกข้อมูลการซื้อเข้าออกจากข้อมูลปริมาณคงเหลือ เพราะใบเสร็จเพียงอย่างเดียวไม่บอกอัตราการใช้จริง [12]")
subheading("4.4", "การตรวจจับเอกสารและเครื่องมือส่งออก")
p("mAP@0.5 ใช้วัดการตรวจจับที่เกณฑ์ IoU 0.5 ส่วน mAP@0.5:0.95 ใช้หลายเกณฑ์ [6] สำหรับภาษาไทย Tesseract มีข้อมูลภาษา tha และ EasyOCR รองรับ th ร่วมกับ en [7, 8] Rockchip แยก RKNN Toolkit2 สำหรับการแปลง/ทดสอบโมเดลภาพและ RKLLM สำหรับการแปลง/รันโมเดลภาษา จึงต้องตรวจความเข้ากันได้ของแต่ละส่วนแยกกัน [9, 10]")

heading("5", "วิธีดำเนินการวิจัย")
subheading("5.1", "สถาปัตยกรรมรวมและผลนำร่องที่มีแล้ว")
p("ภาพที่ 1 แสดง e-Senses ของบ้านและเส้นทางศึกษาระยะแรก ชั้นรับรู้ใช้ e-Eye รับภาพใบซื้อสินค้า ตรวจกรอบเอกสารด้วย YOLO และอ่านด้วย OCR ชั้นประมวลผลจะใช้โมเดลภาษาขนาดเล็กสกัด JSON พร้อมตรวจ schema และความสมเหตุสมผล ก่อนส่งต่อชั้นประยุกต์ใช้ผ่าน MQTT/Home Assistant เมื่อข้อมูลพร้อม ส่วน e-Nose และ e-Skin ยังเป็นโหนดแนวคิดสำหรับระยะต่อเนื่อง [12]")
figure = doc.add_paragraph()
figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
figure.paragraph_format.keep_with_next = True
figure.add_run().add_picture(str(ARCH), width=Inches(5.65))
caption("ภาพ", 1, "กรอบสถาปัตยกรรม e-Senses และเส้นทางศึกษา e-Eye ระยะแรก (ออกแบบโดยผู้วิจัย)")
p("ปัจจุบันฝึก YOLOv8n ด้วย PC CPU และตรวจ OCR แบบ spot-check เท่านั้น RKNN Toolkit2 บน PC แปลง ONNX เป็น RKNN เป้าหมาย RK3588 แบบไม่ quantize และรัน inference ได้หนึ่งภาพ ยังไม่มีการประเมิน JSON/LLM แบบมีเฉลยตรวจซ้ำหรือการวัดบนบอร์ด การตรวจเพิ่มพบว่าภาพส่งเพิ่มทั้ง 6 ใบมีตัวหนังสือหมุนข้าง เมื่อลองหมุนภาพและกรอบเฉลยให้ตรงโดยไม่ฝึกเพิ่ม mAP@0.5 ของ 6 ภาพเปลี่ยนจาก 0.237 เป็น 0.714 อีกด้านหนึ่งการหมุน 7 ภาพที่อ่านตรงในชุดแรกไป 90 องศาทำให้คะแนนเปลี่ยนจาก 0.995 เป็น 0.492 ผลแบบจับคู่นี้ชี้ความไวต่อการหมุนเฉพาะชุดเล็ก ยังไม่พิสูจน์สาเหตุเดียวของคะแนนต่างระหว่างชุด [4, 9]")
caption("ตาราง", 1, "ผลนำร่องและขอบเขตการตีความ")
table(["รายการ", "ผล ณ 17 ต.ค. 2569", "ข้อจำกัด"], [
    ("ข้อมูล", "124 ไฟล์ / 113 ไม่ซ้ำ / 59 ผู้สมัคร", "ภาพไม่เท่ากับธุรกรรมอิสระ"),
    ("YOLOv8n PC", "38 ฝึก / 10 ตรวจ; mAP@0.5 0.928", "ฉากเดิมคล้ายกัน; label รอตรวจซ้ำ"),
    ("ชุดต่างฉาก", "6 ภาพ; mAP@0.5 0.237", "กลุ่มเล็ก; ยังไม่ใช่ final test"),
    ("EasyOCR", "3 ภาพ; spot-check วันที่/ยอด 4/6", "ตรวจคนรอบเดียว; ไม่ใช่ CER"),
    ("RKNN PC", "ONNX → RKNN → PC inference สำเร็จ", "ไม่ใช่ benchmark NPU จริง"),
], [1.35, 2.05, 2.3])

subheading("5.2", "ชุดข้อมูลและเฉลยอ้างอิง")
p("ใช้เฉพาะใบซื้อสินค้าที่มองเห็นรายการสินค้า จำนวน และราคาเป็นชุดหลัก เอกสารชำระเงินหรือใบค่าหอเก็บเป็นชุดเสริมหรือทดสอบชนิดเอกสาร ไม่ปะปนกับคะแนนสกัดรายการสินค้า 59 ภาพปัจจุบันเป็นผู้สมัคร; จะรวมภาพที่ถ่ายซ้ำและเอกสารคู่เป็นกลุ่มธุรกรรมก่อนแบ่งชุดฝึก/ตรวจ/ทดสอบ เพื่อไม่ให้ข้อมูลใกล้ซ้ำข้ามชุด เป้าหมายระยะถัดไปคือเพิ่มความหลากหลายของร้าน ฉาก แสง มุม และอุปกรณ์ถ่าย และขยายเป็นอย่างน้อย 100 ธุรกรรมอิสระหากการเก็บข้อมูลและเวลาทำได้")
p("การตรวจชื่อสินค้าเบื้องต้นในภาพผู้สมัคร 59 ภาพพบว่า ถอดชื่อครบครั้งแรก 12 ภาพ อ่านได้บางส่วนแต่ยังไม่ครบ 45 ภาพ และอีก 2 ภาพไม่มีบรรทัดชื่อสินค้า ตัวเลขอักษรจากส่วนที่อ่านได้ยังใช้แทนสัดส่วนภาษาของทั้งชุดไม่ได้ และยังไม่มีผู้ตรวจซ้ำ")
p("เฉลยแต่ละภาพจะประกอบด้วยกรอบเอกสาร ข้อความอ้างอิง และ JSON schema คงที่: transaction_date, merchant, items[{name, quantity, unit, unit_price, line_total}], grand_total, currency และช่องสถานะอ่านไม่ออก กำหนด null สำหรับข้อมูลไม่ปรากฏ ไม่เดาจากบริบท จะให้ผู้วิจัยสองคนตรวจเฉลยแยกกันแล้วแก้ข้อขัดแย้งร่วมกัน พร้อมบันทึก version, กลุ่มธุรกรรม และเหตุผลการตัดภาพออก การวัดมุมเอียงละเอียดต้องกำหนดแนวอ้างอิงของกระดาษหรือบรรทัด ส่วนสัดส่วนตัวอักษรไทย/อังกฤษจะนับจากข้อความเฉลยที่ตรวจแล้ว ไม่ใช้ OCR ดิบเป็นค่าจริง")

subheading("5.3", "การทดลองและตัวชี้วัด")
p("การตรวจจับใช้กรอบ receipt_region หนึ่งบริเวณต่อภาพเท่าที่นิยามอนุญาต รายงาน mAP@0.5, mAP@0.5:0.95, precision, recall และสัดส่วนกรอบคะแนนสูงสุดที่ IoU≥0.5 แยกฉากเดิมกับฉากใหม่ พร้อมรายงานจำนวนภาพและช่วงความไม่แน่นอนเมื่อขนาดชุดทดสอบเหมาะสม การตรวจ OCR เปรียบเทียบอย่างน้อย EasyOCR กับ Tesseract บนภาพและการ crop เดียวกัน ใช้ CER หลังประกาศกฎ normalize Unicode/ช่องว่าง และวัดความถูกต้องวันที่ ยอดรวม และบรรทัดสินค้าแยกจาก CER [6–8]")
p("การสกัด JSON ใช้ข้อความ OCR เดียวกันในแต่ละเงื่อนไข และใช้ข้อความเฉลยแทน OCR อีกชุดหนึ่งเพื่อแยกความผิดพลาดต้นทาง เปรียบเทียบอย่างน้อยสองโมเดลขนาดเล็กที่ตรวจสิทธิและความเข้ากันได้แล้ว × สองระดับ quantization × เปิด/ปิด constrained decoding ภายใต้ prompt, schema, seed และขีดจำกัดโทเคนเดียวกัน ตัวชี้วัดหลักคือ parse rate, schema-valid rate, exact-match รายฟิลด์, ความคลาดเคลื่อนจำนวนเงิน และความครบของรายการสินค้า ตัวชี้วัดทรัพยากรคือเวลา end-to-end, เวลาแต่ละขั้น, peak RAM และขนาดโมเดล รายงานผลตามธุรกรรมและใช้การสุ่มซ้ำแบบจัดกลุ่มธุรกรรมเพื่อประมาณช่วงความไม่แน่นอน ไม่ตีความความต่างเล็กน้อยจากภาพใกล้ซ้ำว่าเป็นหลักฐานอิสระ")
caption("ตาราง", 2, "จุดตัดสินใจของการทดลอง")
table(["ขั้น", "เปรียบเทียบ", "เกณฑ์ตัดสินหลัก"], [
    ("ตรวจจับ", "ฉากเดิม / ต่างฉาก", "mAP และกรอบสูงสุด IoU≥0.5"),
    ("OCR", "EasyOCR / Tesseract", "CER และความถูกต้องของฟิลด์"),
    ("LLM", "โมเดล × quantization × grammar", "schema validity / field accuracy / RAM / latency"),
    ("ส่งออก", "ONNX / RKNN PC; RKLLM เมื่อพร้อม", "ผ่าน conversion และผลเลขสมเหตุสมผล"),
], [1.0, 2.1, 2.6])

subheading("5.4", "การจำลองบน PC และการตรวจบนอุปกรณ์ในอนาคต")
p("ระยะปัจจุบันใช้ PC เป็นฐานวัดและใช้ SDK เพื่อตรวจการแปลงโมเดลสำหรับเป้าหมาย RK3588 โดยบันทึกเวอร์ชันเครื่องมือ โมเดล ค่าพารามิเตอร์ และข้อมูลทดสอบ การจำกัด thread/RAM บน PC เป็นเพียงการศึกษาความไวต่อทรัพยากร ไม่ใช่การจำลองเวลา NPU แบบ cycle-accurate และไม่ใช้แทนการวัดพลังงานจริง [9, 10]")
p("สเปกอ้างอิงสำหรับการขอยืมภายหลังคือ Radxa ROCK 5B หรือบอร์ดเทียบเท่าที่ใช้ RK3588, RAM 16 GB, ที่เก็บข้อมูลอย่างน้อย 64 GB, Linux image ที่มี NPU driver และ RKNN/RKLLM runtime ที่เข้ากันได้ พร้อมไฟเลี้ยงและการระบายความร้อน หากคณะมีรุ่น 8 GB จะประเมินขนาดโมเดลอีกครั้ง สเปกนี้เป็นกรอบคัดเลือก ไม่ใช่รายการจัดซื้อหรือผลทดสอบ อุปกรณ์จริงต้องยืนยันเวอร์ชัน SDK และความสามารถ constrained decoding ก่อนกำหนดวิธีทดลองบน NPU [9–11, 13]")

subheading("5.5", "การเชื่อมต่อปลายน้ำและความเป็นส่วนตัว")
p("เมื่อ JSON ผ่านการตรวจแล้วจึงทดสอบส่งข้อความจำลองผ่าน MQTT ไปยังระบบบ้านอัจฉริยะ โดยแยกคะแนนความถูกต้องของข้อมูลต้นทางออกจากการส่งข้อความ ใบเสร็จบอกการซื้อเข้า ไม่ได้บอกอัตราการบริโภคโดยตรง การพยากรณ์ burn rate จึงต้องมีข้อมูลปริมาณคงเหลือหรือการใช้จริงตามเวลาเพิ่มเติม และยังไม่ใช่ข้ออ้างผลสำเร็จของระยะนี้")
p("เก็บภาพดิบและ OCR ดิบในที่เก็บส่วนตัว จำกัดผู้เข้าถึง ใช้รหัสภาพแทนชื่อ และลบหรือปกปิดชื่อ เบอร์โทร เลขสมาชิก ที่อยู่ และข้อมูลการชำระเงินก่อนเผยแพร่ผล หากจะเก็บข้อมูลจากผู้เข้าร่วมหรือประเมินอัตราการแก้ไขโดยมนุษย์ ให้หารืออาจารย์และกระบวนการจริยธรรมของคณะก่อนเริ่มขั้นนั้น")

heading("6", "ขอบเขตและข้อจำกัดของการวิจัย")
p("ขอบเขตเชิงสถาปัตยกรรมครอบคลุม e-Senses และกระบวนการสามชั้นตามภาพที่ 1 ขอบเขตการทดลองหลักระยะแรกใช้ภาพใบซื้อสินค้าภาษาไทย/อังกฤษเพื่อตรวจจับ OCR และสกัด JSON บน PC พร้อมตรวจความเข้ากันได้ของ RKNN/RKLLM ตามความพร้อมของเครื่องมือ e-Nose, e-Skin, กล้อง IP แบบเรียลไทม์ และระบบบ้านอัจฉริยะเต็มรูปแบบเป็นส่วนต่อยอด ยังไม่มีบอร์ดสำหรับทดสอบ RK3588 จริง จึงไม่รวมผล NPU ในผลนำร่อง")
p("ข้อจำกัดปัจจุบันคือข้อมูลหลายภาพอาจมาจากธุรกรรมเดียวกัน label ตรวจโดยผู้ช่วยและต้องให้ผู้วิจัยตรวจซ้ำ จำนวนทดสอบต่างฉากยังเล็ก และเฉลยข้อความ/JSON ยังไม่ครบ จึงไม่รายงาน CER, field accuracy หรือข้อสรุปด้าน NPU จากหลักฐานปัจจุบัน")

heading("7", "สถานที่ทำวิจัย")
p("ดำเนินการเตรียมข้อมูล ฝึกและทดสอบเบื้องต้นบนคอมพิวเตอร์ส่วนตัวของผู้จัดทำและสภาพแวดล้อม PC/WSL ที่บันทึกเวอร์ชันไว้ ใช้พื้นที่เก็บข้อมูลส่วนตัวที่ไม่ซิงก์ภาพดิบขึ้นโฟลเดอร์ส่งงาน สถานที่ทดสอบบอร์ดจริงจะกำหนดภายหลังร่วมกับอาจารย์ที่ปรึกษาเมื่อมีอุปกรณ์")

heading("8", "ประโยชน์ที่คาดว่าจะได้รับ")
bullet("ได้กรอบสถาปัตยกรรม e-Senses และข้อกำหนดการส่งข้อมูลระหว่างชั้นที่ใช้พัฒนาระบบบ้านอัจฉริยะในระยะต่อไป")
bullet("ได้ชุดข้อมูลและเฉลยที่ตรวจสอบย้อนหลังได้สำหรับใบซื้อสินค้าไทย/อังกฤษ พร้อมข้อกำหนดการแบ่งชุดตามธุรกรรม")
bullet("ได้หลักฐานว่าแต่ละขั้นของ e-Eye ผิดพลาดตรงไหน โดยเฉพาะเมื่อฉากภาพเปลี่ยน และมีวิธีประเมิน JSON ที่แยกความถูกต้องโครงสร้างจากเนื้อหา")
bullet("ได้ข้อมูล trade-off บน PC ที่ช่วยคัดแบบจำลองและกำหนดการทดลองบนฮาร์ดแวร์จริงอย่างมีเหตุผล โดยไม่สับสนผลจำลองกับ NPU")

heading("9", "แผนและระยะเวลาดำเนินการ")
caption("ตาราง", 3, "แผนดำเนินงานระยะ PC นับจากการรับรองเค้าโครง")
table(["ช่วง", "กิจกรรม", "ผลส่งมอบ"], [
    ("สัปดาห์ 1", "ตรวจประเภทภาพ/กลุ่มธุรกรรมและ label", "manifest และ split ที่ตรวจซ้ำ"),
    ("สัปดาห์ 2–3", "ทำเฉลยข้อความและ JSON สองรอบ", "gold set พร้อมคู่มือ annotation"),
    ("สัปดาห์ 3–4", "เปรียบเทียบ OCR และปรับตรวจจับข้ามฉาก", "CER / field metrics / detection metrics"),
    ("สัปดาห์ 4–5", "ทดลองโมเดลภาษาและ constrained output บน PC", "ตารางความแม่นยำ เวลา และ RAM"),
    ("สัปดาห์ 6", "ทวน RKNN/RKLLM PC-side และเขียนผล", "รายงานข้อจำกัดและขั้นทดสอบบอร์ด"),
], [0.95, 2.75, 2.0])
p("หลังระยะ PC หากมีบอร์ดยืมและการอนุมัติที่เกี่ยวข้อง จึงย้ายโมเดลที่เข้ากันได้ไปวัด latency, RAM และพลังงานจริงบนอุปกรณ์ โดยคงชุดทดสอบและเวอร์ชันโมเดลไว้เทียบกับ PC กิจกรรมนี้เป็นระยะต่อเนื่องภายหลังการศึกษาเบื้องต้น")

heading("10", "งบประมาณ")
p("ระยะ PC ใช้คอมพิวเตอร์ พื้นที่เก็บข้อมูล และซอฟต์แวร์ที่มีสิทธิใช้อยู่ จึงไม่มีวงเงินจัดซื้อฮาร์ดแวร์ในระยะนี้ ส่วนบอร์ด RK3588 และอุปกรณ์ประกอบจะดำเนินเรื่องขอยืมผ่านคณะเมื่อเริ่มระยะทดสอบจริง หากต้องจัดหาใหม่ จะเสนอรายการรุ่น ราคา และความเข้ากันได้ต่ออาจารย์ก่อนกำหนดงบประมาณ ไม่ใส่ราคาที่ยังไม่ได้สำรวจ")

refs_heading = heading("11", "เอกสารอ้างอิง")
refs_heading.paragraph_format.page_break_before = True
refs = [
    "[1] Husom, E. J., et al. (2025). Sustainable LLM inference for edge AI: Evaluating quantized LLMs for energy efficiency, output accuracy, and inference latency. ACM Transactions on Internet of Things, 6(4), Article 28. https://doi.org/10.1145/3767742",
    "[2] Giedra, H., et al. (2026). Benchmarking large language model inference on limited-resource edge systems. Electronics, 15(11), 2451. doi:10.3390/electronics15112451.",
    "[3] Zheng, L., et al. (2024). SGLang: Efficient execution of structured language model programs. Advances in Neural Information Processing Systems, 37, 62557-62583. https://doi.org/10.52202/079017-2000",
    "[4] ณัชพล เพ็งพล และ ธนบดี เริ่มรักสกุล. (2569). ผลนำร่อง e-Eye บน PC ณ วันที่ 17 ตุลาคม 2569. บันทึกผลภายในโครงงาน.",
    "[5] ggml-org. (2026). GBNF guide, llama.cpp. https://github.com/ggml-org/llama.cpp/tree/master/grammars",
    "[6] Ultralytics. (2026). YOLO performance metrics. https://docs.ultralytics.com/guides/yolo-performance-metrics/",
    "[7] Tesseract OCR. (2026). Data files and language support. https://github.com/tesseract-ocr/tessdata",
    "[8] Jaided AI. (2026). EasyOCR. https://github.com/JaidedAI/EasyOCR",
    "[9] Rockchip. (2026). RKNN Toolkit2. https://github.com/airockchip/rknn-toolkit2",
    "[10] Rockchip. (2026). RKLLM. https://github.com/airockchip/rknn-llm",
    "[11] Rockchip. (2022). RK3588 brief datasheet. https://www.rock-chips.com/a/en/products/RK35_Series/2022/0926/1660.html",
    "[12] ณัชพล เพ็งพล และ ธนบดี เริ่มรักสกุล. (2569). การพัฒนาระบบประสาทส่วนกลางอัจฉริยะในที่พักอาศัยเพื่อสกัดและวิเคราะห์ข้อมูลห่วงโซ่อุปทานในครัวเรือนด้วยปัญญาประดิษฐ์ประมวลผลเฉพาะถิ่น. รายงานทบทวนวรรณกรรม 32 หน้า, เอกสารภายในโครงงาน.",
    "[13] Radxa. (2026). ROCK 5B introduction and specifications. https://docs.radxa.com/en/rock5/rock5b/getting-started/introduction",
    "[14] Rivkin, D., et al. (2025). AIoT smart home via autonomous LLM agents. IEEE Internet of Things Journal, 12(3), 2458-2472. https://doi.org/10.1109/JIOT.2024.3471904",
    "[15] Sawalha, S., & Al-Naymat, G. (2025). An adaptive Holt-Winters model for seasonal forecasting of Internet of Things (IoT) data streams. IoT, 6(3), Article 39. https://doi.org/10.3390/iot6030039",
    "[16] Serepas, F., et al. (2025). Lightweight embedded IoT gateway for smart homes based on an ESP32 microcontroller. Computers, 14(9), Article 391. https://doi.org/10.3390/computers14090391",
    "[17] Yin, W., et al. (2026). ShadowNPU: System and algorithm co-design for NPU-centric on-device LLM inference. Proceedings of MobiSys 2026, 294-308. https://doi.org/10.1145/3745756.3809205",
]
for ref in refs:
    para = p(ref)
    para.paragraph_format.left_indent = Inches(0.25)
    para.paragraph_format.first_line_indent = Inches(-0.25)
    para.paragraph_format.space_after = Pt(2)

p("ลงชื่อ ................................................ (นายณัชพล เพ็งพล)  ผู้เสนอโครงงาน", align=WD_ALIGN_PARAGRAPH.RIGHT)
p("ลงชื่อ ................................................ (นายธนบดี เริ่มรักสกุล)  ผู้เสนอโครงงาน", align=WD_ALIGN_PARAGRAPH.RIGHT)
p("วันที่ 17 ตุลาคม 2569", align=WD_ALIGN_PARAGRAPH.RIGHT)

doc.core_properties.title = "เค้าโครงโครงงานระบบประสาทส่วนกลางอัจฉริยะในที่พักอาศัย"
doc.core_properties.subject = "Computer Project Proposal"
doc.core_properties.author = "ณัชพล เพ็งพล และ ธนบดี เริ่มรักสกุล"
doc.save(OUT)
print(OUT)
